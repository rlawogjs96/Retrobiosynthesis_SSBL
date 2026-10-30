import ast
import csv
import datetime
import json
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

# Optional dependency (installed via pip) used for Excel output
from openpyxl import Workbook

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_TSV = SCRIPT_DIR / "inhouse_antismash_t1pks_genes.tsv"


###############################################################################
# Data structures
###############################################################################


@dataclass(frozen=True)
class DomainPattern:
    """Pattern for a single domain in the user-specified module."""

    type: str  # e.g. "KS", "AT", "KR"
    optional: bool = False  # Only meaningful for KS in the current use case
    substrates: Optional[Set[str]] = None  # Allowed substrates for AT (normalized)
    kr_type: Optional[str] = None  # KR type (e.g., "B1", "B2") if specified


@dataclass(frozen=True)
class ModulePattern:
    """Pattern for a module in the user-specified sequence."""

    domains: Tuple[DomainPattern, ...]


@dataclass(frozen=True)
class DomainInstance:
    """Domain instance from a gene module in the TSV file."""

    type: str
    substrate: Optional[str] = None  # Normalized substrate for AT
    kr_type: Optional[str] = None  # KR type (e.g., "B1", "B2") if present


@dataclass(frozen=True)
class GeneModule:
    """One PKS/NRPS module from a gene in the TSV file."""

    domains: Tuple[DomainInstance, ...]


@dataclass(frozen=True)
class GeneRecord:
    """Gene record with ordered modules."""

    gene_id: str  # bgc_id::locus_tag
    bgc_id: str
    locus_tag: str
    modules: Tuple[GeneModule, ...]


@dataclass(frozen=True)
class GeneSegment:
    """
    A contiguous slice of a gene's modules.

    This lets us optionally "split" a gene into smaller segments when matching.
    """

    gene: GeneRecord
    module_start: int  # inclusive index in gene.modules
    module_end: int  # exclusive index in gene.modules
    modules: Tuple[GeneModule, ...]


###############################################################################
# Normalization helpers
###############################################################################


def _norm_text(s: str) -> str:
    """Normalize free-text fields such as substrates."""

    s = s.lower()
    cleaned_chars = []
    for ch in s:
        if ch.isalnum():
            cleaned_chars.append(ch)
        else:
            cleaned_chars.append(" ")
    out = "".join(cleaned_chars)
    # Collapse whitespace
    out = " ".join(out.split())
    return out


###############################################################################
# Parsing user-specified module strings
###############################################################################


def parse_module_string(module_str: str) -> ModulePattern:
    """
    Parse a human-readable module string such as:

        "(KS) - AT(acetyl coa or malonyl coa)"
        "KS - AT(malonyl coa) - KR - DH"
        "KR(B1, active)" or "KR(B2)"

    into a ModulePattern.
    """

    parts = [p.strip() for p in module_str.split("-") if p.strip()]
    domains: List[DomainPattern] = []

    for token in parts:
        optional = False
        # Optional KS syntax: "(KS)"
        if token.startswith("(") and token.endswith(")"):
            inner = token[1:-1].strip()
            token = inner
            optional = True

        # Handle AT with substrates, e.g. "AT(acetyl coa or malonyl coa)"
        if token.upper().startswith("AT"):
            # Extract substrates inside parentheses if present
            substrates: Optional[Set[str]] = None
            if "(" in token and token.endswith(")"):
                idx = token.index("(")
                sub_str = token[idx + 1 : -1]
                # Split on "or"
                raw_subs = [s.strip() for s in sub_str.split("or") if s.strip()]
                substrates = {_norm_text(s) for s in raw_subs}
            domain_type = "AT"
            domains.append(DomainPattern(type=domain_type, optional=optional, substrates=substrates, kr_type=None))
        elif token.upper().startswith("KR"):
            # Handle KR with type, e.g. "KR(B1, active)" or "KR(B2)"
            domain_type = "KR"
            kr_type: Optional[str] = None
            if "(" in token and token.endswith(")"):
                idx = token.index("(")
                content = token[idx + 1 : -1].strip()
                # Extract the first part before comma (e.g., "B1" from "B1, active")
                if "," in content:
                    kr_type = content.split(",")[0].strip().upper()
                else:
                    kr_type = content.strip().upper()
                # If empty after processing, set to None
                if not kr_type:
                    kr_type = None
            domains.append(DomainPattern(type=domain_type, optional=optional, substrates=None, kr_type=kr_type))
        else:
            # Plain domain, e.g. "KS", "DH", "ER", "TE"
            domain_type = token.upper()
            domains.append(DomainPattern(type=domain_type, optional=optional, substrates=None, kr_type=None))

    return ModulePattern(domains=tuple(domains))


def parse_module_list(module_strings: List[str]) -> List[ModulePattern]:
    """Parse a list of human-readable module strings into patterns."""

    return [parse_module_string(s) for s in module_strings]


###############################################################################
# Parsing domain strings from TSV format
###############################################################################


def _parse_domain_string(domain_str: str) -> Optional[DomainInstance]:
    """
    Parse a domain string from TSV format into a DomainInstance.
    
    Examples:
        "KS" -> DomainInstance(type="KS", substrate=None, kr_type=None)
        "AT(Malonyl-CoA)" -> DomainInstance(type="AT", substrate="malonyl coa", kr_type=None)
        "KR(B1, active)" -> DomainInstance(type="KR", substrate=None, kr_type="B1")
        "KR(B2)" -> DomainInstance(type="KR", substrate=None, kr_type="B2")
        "PP" -> DomainInstance(type="PP", substrate=None, kr_type=None)
        "ACP" -> DomainInstance(type="ACP", substrate=None, kr_type=None)
        "TE(None)" -> DomainInstance(type="TE", substrate=None, kr_type=None)
    """
    
    domain_str = domain_str.strip()
    if not domain_str:
        return None
    
    # Check if it has parentheses (e.g., "AT(Malonyl-CoA)", "KR(B1, active)", "TE(None)")
    if "(" in domain_str and domain_str.endswith(")"):
        idx = domain_str.index("(")
        domain_type = domain_str[:idx].upper()
        content = domain_str[idx + 1 : -1].strip()
        
        substrate_norm: Optional[str] = None
        kr_type: Optional[str] = None
        
        if domain_type == "AT":
            # Extract substrate from content (e.g., "Malonyl-CoA" -> "malonyl coa")
            substrate_norm = _norm_text(content)
        elif domain_type == "KR":
            # Extract KR type from content (e.g., "B1" from "B1, active" or "B2" from "B2")
            if "," in content:
                kr_type = content.split(",")[0].strip().upper()
            else:
                kr_type = content.strip().upper()
            # If empty after processing, set to None
            if not kr_type:
                kr_type = None
        # For other domains like DH, ER, TE, we ignore the content in parentheses
        # (e.g., "TE(None)")
        
        return DomainInstance(type=domain_type, substrate=substrate_norm, kr_type=kr_type)
    else:
        # Plain domain without parentheses
        domain_type = domain_str.upper()
        return DomainInstance(type=domain_type, substrate=None, kr_type=None)


###############################################################################
# Loading and representing genes from the TSV file
###############################################################################


def _build_gene_records(
    tsv_path: str,
    domain_whitelist: Optional[Set[str]] = None,
    verbose: bool = False,
) -> List[GeneRecord]:
    """
    Load in-house antismash TSV and convert to a flat list of GeneRecord objects.

    Parameters
    ----------
    tsv_path:
        Path to inhouse_antismash_t1pks_genes.tsv
    domain_whitelist:
        If provided, only keep these domain types when constructing modules.
        By default includes KS, AT, KR, DH, ER, TE, ACP, PP.
    """

    if domain_whitelist is None:
        domain_whitelist = {"KS", "AT", "KR", "DH", "ER", "TE", "ACP", "PP"}

    if not os.path.exists(tsv_path):
        raise FileNotFoundError(f"TSV file not found: {tsv_path}")

    gene_records: List[GeneRecord] = []
    
    if verbose:
        print("Loading genes from TSV file...")
    row_count = 0

    with open(tsv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f, delimiter="\t")
        for row in reader:
            row_count += 1
            if verbose and row_count % 10000 == 0:
                print(f"  Processed {row_count:,} rows, found {len(gene_records):,} valid genes...")
            bgc_id = row.get("bgc_id", "unknown_bgc")
            locus_tag = row.get("gene", "unknown_gene")
            gene_id = f"{bgc_id}::{locus_tag}"
            
            # Parse the domains column (it's a string representation of a list of lists)
            domains_str = row.get("domains", "[]")
            try:
                modules_list = ast.literal_eval(domains_str)
            except (ValueError, SyntaxError) as e:
                # Skip rows with invalid domain data
                continue
            
            if not isinstance(modules_list, list):
                continue
            
            gene_modules: List[GeneModule] = []
            for module_domains in modules_list:
                if not isinstance(module_domains, list):
                    continue
                
                domains: List[DomainInstance] = []
                for domain_str in module_domains:
                    if not isinstance(domain_str, str):
                        continue
                    
                    domain_inst = _parse_domain_string(domain_str)
                    if domain_inst is None:
                        continue
                    
                    # Filter by whitelist
                    if domain_inst.type not in domain_whitelist:
                        continue
                    
                    domains.append(domain_inst)
                
                if domains:
                    gene_modules.append(GeneModule(domains=tuple(domains)))
            
            if gene_modules:
                gene_records.append(
                    GeneRecord(
                        gene_id=gene_id,
                        bgc_id=bgc_id,
                        locus_tag=locus_tag,
                        modules=tuple(gene_modules),
                    )
                )

    if verbose:
        print(f"  Finished loading: {row_count:,} rows processed, {len(gene_records):,} genes loaded")
    return gene_records


###############################################################################
# Matching logic
###############################################################################


def _domains_match(pattern: DomainPattern, inst: DomainInstance) -> bool:
    """
    Check if a single domain instance matches the pattern (ignoring optionality).
    
    Special handling:
    - "PP" in gene data matches "ACP" in pattern (and vice versa)
    - For KR: if pattern specifies a type, instance must match; if pattern doesn't specify, exclude C, C1, C2 types
    - For KR types: if pattern is just 'A', 'B', or 'C', it matches all types starting with that letter
      (e.g., "KR(B)" matches "KR(B1)", "KR(B2)", etc.)
    """
    
    # Handle ACP/PP equivalence: pattern "ACP" matches gene "PP" or "ACP", and vice versa
    pattern_type = pattern.type
    inst_type = inst.type
    
    if pattern_type == "ACP" and inst_type == "PP":
        return True
    if pattern_type == "PP" and inst_type == "ACP":
        return True
    
    if pattern_type != inst_type:
        return False

    if pattern.type == "AT":
        # If pattern specifies substrates, they must contain the instance substrate (if present)
        if pattern.substrates:
            if inst.substrate is None:
                return False
            return inst.substrate in pattern.substrates
        # No substrate constraint
        return True

    if pattern.type == "KR":
        # If pattern specifies KR type, check matching
        if pattern.kr_type is not None:
            if inst.kr_type is None:
                return False
            # If pattern is just a single letter (A, B, or C), match any type starting with that letter
            if len(pattern.kr_type) == 1 and pattern.kr_type in ('A', 'B', 'C'):
                return inst.kr_type.startswith(pattern.kr_type)
            # Otherwise, require exact match
            return inst.kr_type == pattern.kr_type
        # If pattern doesn't specify KR type, exclude C, C1, C2 types unless instance has no type
        if inst.kr_type is None:
            return True
        # Exclude C, C1, C2 types when pattern doesn't specify
        if inst.kr_type in ('C', 'C1', 'C2'):
            return False
        # Any other KR type (A, A1, A2, B, B1, B2, etc.) is OK
        return True

    # Other domains: match only by type (already checked above)
    return True


def module_pattern_matches_instance(pattern: ModulePattern, inst: GeneModule) -> bool:
    """
    Check if a user-specified ModulePattern matches a GeneModule.

    Special handling:
    - Leading optional KS in the pattern may be present or absent in the gene module.
    """

    p_domains = list(pattern.domains)
    g_domains = list(inst.domains)

    # Handle possible leading optional domain (e.g. (KS))
    if p_domains and p_domains[0].optional:
        # Two possibilities: pattern uses the optional domain, or it doesn't.
        # 1) Use the optional domain
        if len(p_domains) <= len(g_domains):
            ok1 = True
            for pd, gd in zip(p_domains, g_domains):
                if not _domains_match(pd, gd):
                    ok1 = False
                    break
        else:
            ok1 = False

        # 2) Skip the optional domain
        p2 = p_domains[1:]
        if len(p2) <= len(g_domains):
            ok2 = True
            for pd, gd in zip(p2, g_domains):
                if not _domains_match(pd, gd):
                    ok2 = False
                    break
        else:
            ok2 = False

        # Require exact length match in either branch
        ok1 = ok1 and len(p_domains) == len(g_domains)
        ok2 = ok2 and len(p2) == len(g_domains)
        return ok1 or ok2

    # No leading optional domain: require exact length and element-wise match
    if len(p_domains) != len(g_domains):
        return False

    return all(_domains_match(pd, gd) for pd, gd in zip(p_domains, g_domains))


###############################################################################
# Gene combination search
###############################################################################


def _build_position_edges(
    target_patterns: List[ModulePattern],
    genes: List[GeneRecord],
    allow_gene_splitting: bool = False,
    verbose: bool = False,
) -> Dict[int, List[Tuple[int, GeneSegment]]]:
    """
    For each start index in the target pattern sequence, build all possible
    (end_index, gene) edges where the gene's modules match a contiguous slice
    of target_patterns.
    """

    n = len(target_patterns)
    edges: Dict[int, List[Tuple[int, GeneSegment]]] = {i: [] for i in range(n + 1)}
    
    if verbose:
        print(f"Building position edges for {len(genes):,} genes against {n} target modules...")
    gene_count = 0

    for gene in genes:
        gene_count += 1
        if verbose and gene_count % 1000 == 0:
            total_edges = sum(len(edge_list) for edge_list in edges.values())
            print(f"  Processed {gene_count:,}/{len(genes):,} genes, found {total_edges:,} edges so far...")
        g_len = len(gene.modules)
        if g_len == 0:
            continue

        # Decide which contiguous module slices of this gene are allowed
        if allow_gene_splitting:
            segment_ranges: List[Tuple[int, int]] = []
            for seg_start in range(g_len):
                for seg_end in range(seg_start + 1, g_len + 1):
                    segment_ranges.append((seg_start, seg_end))
        else:
            segment_ranges = [(0, g_len)]

        for seg_start, seg_end in segment_ranges:
            seg_modules = gene.modules[seg_start:seg_end]
            seg_len = seg_end - seg_start

            max_start = n - seg_len
            if max_start < 0:
                continue

            for start in range(0, max_start + 1):
                ok = True
                for offset in range(seg_len):
                    tp = target_patterns[start + offset]
                    gm = seg_modules[offset]
                    if not module_pattern_matches_instance(tp, gm):
                        ok = False
                        break
                if ok:
                    end = start + seg_len
                    segment = GeneSegment(
                        gene=gene,
                        module_start=seg_start,
                        module_end=seg_end,
                        modules=seg_modules,
                    )
                    edges[start].append((end, segment))
    
    total_edges = sum(len(edge_list) for edge_list in edges.values())
    if verbose:
        print(f"  Finished building edges: {total_edges:,} total edges found")
    return edges


def _reconstruct_paths(
    parents: Dict[int, List[Tuple[int, GeneSegment]]],
    pos: int,
    max_paths: Optional[int] = None,
) -> List[List[Tuple[int, int, GeneSegment]]]:
    """
    Reconstruct paths from position 0 to 'pos' using the parent lists.

    Each path is a list of (start_index, end_index, GeneSegment).
    Guarantees that EVERY optimal split pattern and EVERY candidate GeneSegment
    for each (start, end) slot is represented before applying `max_paths`,
    so `unique_genes_by_gene_coverage`, `print_gene_coverage_counts`, JSON,
    and Excel outputs never truncate candidate gene counts when the Cartesian
    product exceeds `max_paths`.
    """
    if pos == 0:
        return [[]]

    # 1. Group optimal parent segments by (prev_pos, cur_pos) span
    span_segments: Dict[Tuple[int, int], List[GeneSegment]] = {}
    prev_positions: Dict[int, List[int]] = {}
    for cur_p, entries in parents.items():
        seen_prevs: Set[int] = set()
        for prev_p, seg in entries:
            span_segments.setdefault((prev_p, cur_p), []).append(seg)
            if prev_p not in seen_prevs:
                seen_prevs.add(prev_p)
                prev_positions.setdefault(cur_p, []).append(prev_p)

    # 2. Enumerate all minimal-length split patterns ((0, i_1), ..., (i_{k-1}, pos))
    split_patterns: List[Tuple[Tuple[int, int], ...]] = []
    pat_stack: List[Tuple[int, List[Tuple[int, int]]]] = [(pos, [])]
    while pat_stack:
        cur_p, pat = pat_stack.pop()
        if cur_p == 0:
            split_patterns.append(tuple(pat))
            continue
        for prev_p in prev_positions.get(cur_p, []):
            pat_stack.append((prev_p, [(prev_p, cur_p)] + pat))
    split_patterns.sort()

    # 3. First pass: for every split pattern, emit `max_slot_len` striped paths
    #    so every candidate gene in every slot is 100% covered in `paths`.
    paths: List[List[Tuple[int, int, GeneSegment]]] = []
    seen_path_keys: Set[Tuple[Tuple[int, int, str, int, int], ...]] = set()

    for pat in split_patterns:
        slot_lists = [span_segments[(s, e)] for s, e in pat]
        if any(not sl for sl in slot_lists):
            continue
        max_len = max(len(sl) for sl in slot_lists)
        for i in range(max_len):
            combo = [
                (s, e, slot_lists[idx][i % len(slot_lists[idx])])
                for idx, (s, e) in enumerate(pat)
            ]
            key = tuple(
                (s, e, seg.gene.gene_id, seg.module_start, seg.module_end)
                for s, e, seg in combo
            )
            if key not in seen_path_keys:
                seen_path_keys.add(key)
                paths.append(combo)

    if max_paths is not None and len(paths) >= max_paths:
        return paths

    # 4. Second pass: fill remaining budget up to `max_paths` with full DFS combinations
    stack: List[Tuple[int, List[Tuple[int, int, GeneSegment]]]] = [(pos, [])]
    while stack:
        cur_pos, path = stack.pop()
        if cur_pos == 0:
            key = tuple(
                (s, e, seg.gene.gene_id, seg.module_start, seg.module_end)
                for s, e, seg in path
            )
            if key not in seen_path_keys:
                seen_path_keys.add(key)
                paths.append(path)
                if max_paths is not None and len(paths) >= max_paths:
                    break
            continue

        for prev_pos, segment in parents.get(cur_pos, []):
            stack.append((prev_pos, [(prev_pos, cur_pos, segment)] + path))

        if max_paths is not None and len(paths) >= max_paths:
            break

    return paths


def find_gene_combinations(
    target_module_strings: List[str],
    tsv_path: str,
    max_solutions: Optional[int] = None,
    allow_gene_splitting: bool = False,
    max_paths: Optional[int] = 20000,
    verbose: bool = False,
) -> List[List[Tuple[int, int, GeneSegment]]]:
    """
    Find gene combinations that can reconstitute the input list of modules.

    Parameters
    ----------
    target_module_strings:
        List of human-readable module strings in order, e.g.:
        [
            "(KS) - AT(acetyl coa or malonyl coa) - ACP",
            "KS - AT(methylmalonyl coa) - KR - DH - ACP",
            ...
        ]
    tsv_path:
        Path to inhouse_antismash_t1pks_genes.tsv
    max_solutions:
        If provided, limit the number of returned combinations.
    allow_gene_splitting:
        Allow splitting genes into subsegments when matching.
    max_paths:
        Cap the number of reconstructed paths (combinations) to avoid blow-up.
    verbose:
        If True, print progress; otherwise stay quiet.

    Returns
    -------
    List of combinations; each combination is a list of
    (start_index, end_index, GeneSegment), sorted by increasing number of genes.
    """

    if verbose:
        print("=" * 60)
        print("Starting gene search...")
        print("=" * 60)
    
    target_patterns = parse_module_list(target_module_strings)
    if verbose:
        print(f"Parsed {len(target_patterns)} target module patterns")
    
    genes = _build_gene_records(tsv_path=tsv_path, verbose=verbose)
    if not genes:
        if verbose:
            print("No genes found in TSV file.")
        return []

    edges = _build_position_edges(
        target_patterns,
        genes,
        allow_gene_splitting=allow_gene_splitting,
        verbose=verbose,
    )
    n = len(target_patterns)

    # DP to find minimal number of genes to reach each position
    if verbose:
        print("Running dynamic programming to find minimal gene combinations...")
    INF = 10**9
    best: List[int] = [INF] * (n + 1)
    best[0] = 0
    parents: Dict[int, List[Tuple[int, GeneSegment]]] = {}

    for pos in range(n):
        if verbose and pos > 0 and pos % max(1, n // 10) == 0:
            print(f"  DP progress: position {pos}/{n} ({100*pos//n}%)")
        if best[pos] == INF:
            continue
        for end, segment in edges.get(pos, []):
            cost = best[pos] + 1
            if cost < best[end]:
                best[end] = cost
                parents[end] = [(pos, segment)]
            elif cost == best[end]:
                parents.setdefault(end, []).append((pos, segment))
    
    if verbose:
        print(f"  DP completed: minimal number of genes needed = {best[n] if best[n] != INF else 'impossible'}")

    if best[n] == INF:
        if verbose:
            print("No valid gene combinations found.")
        return []

    if verbose:
        print("Reconstructing all paths...")
    all_paths = _reconstruct_paths(parents, n, max_paths=max_paths)
    if verbose:
        print(f"  Found {len(all_paths):,} total combinations")

    # Sort by number of genes (should all be minimal already) and gene_ids
    if verbose:
        print("Sorting combinations...")
    all_paths.sort(key=lambda path: (len(path), [seg[2].gene.gene_id for seg in path]))

    if max_solutions is not None:
        all_paths = all_paths[:max_solutions]
        if verbose:
            print(f"Limited to top {max_solutions} combinations")

    if verbose:
        print("=" * 60)
        print(f"Search complete: {len(all_paths):,} combination(s) found")
        print("=" * 60)
    return all_paths


###############################################################################
# Pretty-printing helpers
###############################################################################


def _format_gene_module(gm: GeneModule) -> str:
    """Render a GeneModule back into a compact string representation."""

    parts: List[str] = []
    for d in gm.domains:
        if d.type == "AT" and d.substrate is not None:
            parts.append(f"AT({d.substrate})")
        elif d.type == "KR" and d.kr_type is not None:
            parts.append(f"KR({d.kr_type})")
        else:
            parts.append(d.type)
    return " - ".join(parts)


def pretty_print_combinations(
    target_module_strings: List[str],
    combinations: List[List[Tuple[int, int, GeneSegment]]],
) -> None:
    """Print combinations in the style requested, ordered by least number of genes."""

    print("Target modules:")
    for i, m in enumerate(target_module_strings):
        print(f"  [{i}] {m}")
    print()

    if not combinations:
        print("No gene combinations found.")
        return

    for ci, combo in enumerate(combinations, start=1):
        print(f"Combination {ci} (genes: {len(combo)}):")
        for gi, (start, end, segment) in enumerate(combo):
            gene = segment.gene
            print(f"  gene_{gi} ({gene.gene_id}) covers target modules [{start}..{end - 1}] "
                  f"with gene modules [{segment.module_start}..{segment.module_end - 1}]:")
            for idx, gm in enumerate(segment.modules, start=segment.module_start):
                print(f"    module_{idx}: {_format_gene_module(gm)}")
        print()


def write_combinations_to_json(
    target_module_strings: List[str],
    combinations: List[List[Tuple[int, int, GeneSegment]]],
    output_path: str,
) -> None:
    """
    Write search results to a JSON file in the format:
    1. Possible ways for the input query to be split
    2. For each split, genes set for 1st position, 2nd position, ..., last position
    
    Parameters
    ----------
    target_module_strings:
        List of target module strings (for reference)
    combinations:
        List of gene combinations found
    output_path:
        Path to output JSON file
    """
    
    if not combinations:
        print(f"No combinations to write to {output_path}")
        return
    
    # Group combinations by split pattern (sequence of target module ranges)
    split_patterns: Dict[Tuple[Tuple[int, int], ...], List[List[Tuple[int, int, GeneSegment]]]] = {}
    
    for combo in combinations:
        # Create a tuple of (start, end) ranges for this combination
        pattern = tuple((start, end) for start, end, _ in combo)
        if pattern not in split_patterns:
            split_patterns[pattern] = []
        split_patterns[pattern].append(combo)
    
    # Build JSON structure
    result = {
        "target_modules": target_module_strings,
        "num_splits": len(split_patterns),
        "splits": []
    }
    
    # Process each unique split pattern
    for split_id, (pattern, combos_with_pattern) in enumerate(sorted(split_patterns.items()), start=1):
        num_positions = len(pattern)
        
        # For each position, collect all unique genes that can fill it
        position_genes: List[Dict[str, Any]] = []
        
        for pos_idx in range(num_positions):
            target_start, target_end = pattern[pos_idx]
            genes_at_position: List[Dict[str, Any]] = []
            seen_gene_ids: Set[str] = set()
            
            # Collect all unique genes that can fill this position
            for combo in combos_with_pattern:
                start, end, segment = combo[pos_idx]
                gene = segment.gene
                
                # Use gene_id as unique identifier
                if gene.gene_id not in seen_gene_ids:
                    seen_gene_ids.add(gene.gene_id)
                    
                    # Format modules as simple list of strings
                    modules = []
                    for gm in segment.modules:
                        modules.append(_format_gene_module(gm))
                    
                    gene_data = {
                        "gene_id": gene.gene_id,
                        "modules": modules
                    }
                    genes_at_position.append(gene_data)
            
            position_genes.append({
                "position": pos_idx + 1,  # 1-indexed position
                "num_genes": len(genes_at_position),
                "genes": genes_at_position
            })
        
        split_data = {
            "split_id": split_id,
            "num_positions": num_positions,
            "target_module_ranges": [
                {"start": start, "end": end - 1} for start, end in pattern
            ],
            "positions": position_genes
        }
        result["splits"].append(split_data)
    
    # Write JSON file with pretty formatting
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(result, f, indent=2, ensure_ascii=False)
    
    print(f"Results written to {output_path}")


def write_splits_to_excel(
    combinations: List[List[Tuple[int, int, GeneSegment]]],
    output_path: str,
) -> None:
    """
    Write an Excel workbook where each worksheet is one split pattern.

    Each worksheet has N columns where N = number of gene positions in that split.
    Each column contains the gene_id values (unique) that can fill that position.
    """

    if not combinations:
        print(f"No combinations to write to {output_path}")
        return

    # Group combinations by split pattern (sequence of target module ranges)
    split_patterns: Dict[Tuple[Tuple[int, int], ...], List[List[Tuple[int, int, GeneSegment]]]] = {}
    for combo in combinations:
        pattern = tuple((start, end) for start, end, _ in combo)
        split_patterns.setdefault(pattern, []).append(combo)

    wb = Workbook()
    # Remove default sheet
    default_ws = wb.active
    wb.remove(default_ws)

    for split_id, (pattern, combos_with_pattern) in enumerate(sorted(split_patterns.items()), start=1):
        num_positions = len(pattern)
        title = f"split_{split_id}"
        ws = wb.create_sheet(title=title[:31])  # Excel sheet name max length is 31

        # Header row
        for pos_idx in range(num_positions):
            ws.cell(row=1, column=pos_idx + 1, value=f"position_{pos_idx + 1}")

        # Collect gene_ids per position
        gene_ids_by_pos: List[List[str]] = []
        for pos_idx in range(num_positions):
            seen: Set[str] = set()
            gene_ids: List[str] = []
            for combo in combos_with_pattern:
                _, _, segment = combo[pos_idx]
                gid = segment.gene.gene_id
                if gid not in seen:
                    seen.add(gid)
                    gene_ids.append(gid)
            gene_ids.sort()
            gene_ids_by_pos.append(gene_ids)

        max_rows = max((len(lst) for lst in gene_ids_by_pos), default=0)
        for r in range(max_rows):
            for c in range(num_positions):
                val = gene_ids_by_pos[c][r] if r < len(gene_ids_by_pos[c]) else ""
                ws.cell(row=2 + r, column=1 + c, value=val)

    wb.save(output_path)
    print(f"Excel results written to {output_path}")


def _format_module_span(start: int, end_exclusive: int) -> str:
    """Human-readable target-module span. Loading is module 0."""
    last = end_exclusive - 1
    names = []
    for i in range(start, end_exclusive):
        names.append("loading" if i == 0 else f"M{i}")
    if start == last:
        return names[0]
    return f"{names[0]}-{names[-1]}"


def unique_genes_by_gene_coverage(
    combinations: List[List[Tuple[int, int, GeneSegment]]],
) -> List[Tuple[Tuple[Tuple[int, int], ...], List[List[str]]]]:
    """Unique gene_ids per gene-coverage slot, grouped by split pattern.

    A coverage slot is one gene in the combination: a contiguous span of
    target modules (start inclusive, end exclusive), matching Excel columns
    and JSON ``positions``.
    """
    split_patterns: Dict[
        Tuple[Tuple[int, int], ...], List[List[Tuple[int, int, GeneSegment]]]
    ] = {}
    for combo in combinations:
        pattern = tuple((start, end) for start, end, _ in combo)
        split_patterns.setdefault(pattern, []).append(combo)

    out: List[Tuple[Tuple[Tuple[int, int], ...], List[List[str]]]] = []
    for pattern, combos_with_pattern in sorted(split_patterns.items()):
        genes_by_slot: List[List[str]] = []
        for pos_idx in range(len(pattern)):
            seen: Set[str] = set()
            gene_ids: List[str] = []
            for combo in combos_with_pattern:
                gid = combo[pos_idx][2].gene.gene_id
                if gid not in seen:
                    seen.add(gid)
                    gene_ids.append(gid)
            genes_by_slot.append(gene_ids)
        out.append((pattern, genes_by_slot))
    return out


def print_gene_coverage_counts(
    combinations: List[List[Tuple[int, int, GeneSegment]]],
) -> None:
    """Print unique genes per gene arrow, not per individual module."""
    print("Unique genes by gene-coverage position")
    print("(one gene may span several target modules):")
    for split_id, (pattern, genes_by_slot) in enumerate(
        unique_genes_by_gene_coverage(combinations), start=1
    ):
        spans = " | ".join(_format_module_span(s, e) for s, e in pattern)
        print(
            f"  split_{split_id}: {len(pattern)} genes covering [{spans}]"
        )
        for pos_idx, ((start, end), genes) in enumerate(
            zip(pattern, genes_by_slot), start=1
        ):
            span = _format_module_span(start, end)
            print(
                f"    position {pos_idx} ({span}): {len(genes)} genes"
            )


def example_search() -> None:
    """Run Collins #12 against the in-repo antiSMASH TSV."""

    tsv_path = str(DEFAULT_TSV)

    # Collins #12
    target_modules = [
        "(KS) - AT(acetyl coa or malonyl coa) - ACP",
        "KS - AT(malonyl coa) - DH - KR - ACP",
        "KS - AT(malonyl coa) - DH - KR - ACP",
        "KS - AT(methylmalonyl coa) - DH - ER - KR - ACP",
        "KS - AT(methylmalonyl coa) - DH - ER - KR - ACP",
        "KS - AT(methylmalonyl coa) - DH - ER - KR - ACP",
        "KS - AT(malonyl coa) - KR - ACP - TE",
    ]

# Collins #4
        # "(KS) - AT(propionyl coa or methylmalonyl coa) - ACP",
        # "KS - AT(methylmalonyl coa) - KR(A2) - ACP",
        # "KS - AT(malonyl coa) - KR(A) - ACP",
        # "KS - AT(malonyl coa) - DH - ER - KR - ACP",
        # "KS - AT(malonyl coa) - DH - KR - ACP",
        # "KS - AT(methylmalonyl coa) - KR(A2) - ACP",
        # "KS - AT(methylmalonyl coa) - DH - KR - ACP",
        # "KS - AT(methylmalonyl coa) - (KR(C)) - ACP - TE",

# Collins #6
        # "(KS) - AT(propionyl coa or methylmalonyl coa) - ACP",
        # "KS - AT(methylmalonyl coa) - DH - ER - KR - ACP",
        # "KS - AT(methylmalonyl coa) - DH - KR - ACP",
        # "KS - AT(malonyl coa) - KR(A) - ACP",
        # "KS - AT(malonyl coa) - KR(B) - ACP",
        # "KS - AT(malonyl coa) - DH - ER - KR - ACP",
        # "KS - AT(malonyl coa) - DH - ER - KR - ACP",
        # "KS - AT(malonyl coa) - DH - ER - KR - ACP - TE",

# Collins #8
        # "(KS) - AT(propionyl coa or methylmalonyl coa) - ACP",
        # "KS - AT(methylmalonyl coa) - KR(B2) - ACP",
        # "KS - AT(methylmalonyl coa) - DH - ER - KR - ACP",
        # "KS - AT(methylmalonyl coa) - KR(B2) - ACP",
        # "KS - AT(malonyl coa) - DH - KR - ACP",
        # "KS - AT(methylmalonyl coa) - (KR(C)) - ACP - TE",

# Collins #12
        # "(KS) - AT(acetyl coa or malonyl coa) - ACP",
        # "KS - AT(malonyl coa) - DH - KR - ACP",
        # "KS - AT(malonyl coa) - DH - KR - ACP",
        # "KS - AT(methylmalonyl coa) - DH - ER - KR - ACP",
        # "KS - AT(methylmalonyl coa) - DH - ER - KR - ACP",
        # "KS - AT(methylmalonyl coa) - DH - ER - KR - ACP",
        # "KS - AT(malonyl coa) - KR - ACP - TE",

# Collins #16
        # "(KS) - AT(propionyl coa or methylmalonyl coa) - ACP",
        # "KS - AT(malonyl coa) - KR(B) - ACP",
        # "KS - AT(methylmalonyl coa) - DH - KR - ACP",
        # "KS - AT(methylmalonyl coa) - KR(A1) - ACP",
        # "KS - AT(malonyl coa) - DH - KR - ACP",
        # "KS - AT(methylmalonyl coa) - KR(A1) - ACP",
        # "KS - AT(malonyl coa) - DH - ER - KR - ACP",
        # "KS - AT(methylmalonyl coa) - (KR(C)) - ACP - TE",

# Collins #17
        # "(KS) - AT(acetyl coa or malonyl coa) - ACP",
        # "KS - AT(malonyl coa) - KR(B) - ACP",
        # "KS - AT(malonyl coa) - DH - ER - KR - ACP",
        # "KS - AT(methylmalonyl coa) - DH - ER - KR - ACP",
        # "KS - AT(malonyl coa) - DH - ER - KR - ACP",
        # "KS - AT(methylmalonyl coa) - KR(B1) - ACP",
        # "KS - AT(methylmalonyl coa) - KR(A1) - ACP - TE",
        # "KS - AT(malonyl coa) - (KR(C)) - ACP - TE",


# Collins #20
        # "(KS) - AT(acetyl coa or malonyl coa) - ACP",
        # "KS - AT(methylmalonyl coa) - DH - KR - ACP",
        # "KS - AT(malonyl coa) - KR(B) - ACP",
        # "KS - AT(methylmalonyl coa) - DH - ER - KR - ACP",
        # "KS - AT(malonyl coa) - KR(B) - ACP",
        # "KS - AT(methylmalonyl coa) - DH - ER - KR - ACP - TE",

    print(f"TSV: {tsv_path}")
    combos = find_gene_combinations(
        target_modules,
        tsv_path=tsv_path,
        max_solutions=None,
        allow_gene_splitting=False,
        verbose=True,
    )

    print_gene_coverage_counts(combos)

    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    output_json = str(SCRIPT_DIR / f"gene_search_results_{timestamp}.json")
    write_combinations_to_json(target_modules, combos, output_json)

    output_xlsx = str(SCRIPT_DIR / f"gene_search_results_{timestamp}.xlsx")
    write_splits_to_excel(combos, output_xlsx)


if __name__ == "__main__":
    # Simple CLI: run the hard-coded example.
    example_search()

